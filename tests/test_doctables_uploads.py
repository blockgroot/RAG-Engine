"""Charts from files uploaded to a chat (`app/doctables/uploads.py`).

Pinned here:

- each format reads the tables it really has: CSV, every Excel sheet, Word
  tables, the tables a PDF page draws, pipe tables in markdown;
- an upload's tables are its uploader's, in that chat, and nobody else's --
  not another member, not the same person in another chat, and never through
  ``list_tables`` (the document listing);
- deleting the file takes its tables with it;
- a file with no table is read for figures ONCE, on demand, in Chart mode;
- Chart mode offers the chat's files first, and a pick runs end to end.
"""

from __future__ import annotations

import io
import uuid

import pytest

from app.agent import insights_agent, routing
from app.attachments import store as attach_store
from app.attachments.extract import extract_text
from app.db import get_connection
from app.doctables import store as table_store, uploads
from app.doctables.store import UploadScope
from app.insights import tables as doc_tables
from app.insights.resolve import ChartSpec, TABLE_METRIC
from app.vectorstore.base import Viewer
from .conftest import requires_db
from .test_attachment_storage import FakeBlobStore

SALES_CSV = "Region,Revenue\nNorth,120\nSouth,90\nNorth,30\n"


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sales = book.active
    sales.title = "Sales"
    for row in [("Region", "Revenue"), ("North", 120), ("South", 90.0), (None, None),
                ("North", 30)]:
        sales.append(row)
    costs = book.create_sheet("Costs")
    for row in [("Team", "Cost"), ("Core", 10), ("Growth", 5)]:
        costs.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Quarterly budget")
    table = document.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Team", "Budget"), ("Core", "10"), ("Growth", "5")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def ruled_pdf(rows: list[list[str]]) -> bytes:
    """A one-page PDF that DRAWS a table (ruled lines + text), as an exported
    report does. Hand-written so the test needs no PDF writer."""
    x0, y_top, cw, rh = 50, 750, 120, 20
    ops = ["0.5 w"]
    n, m = len(rows), len(rows[0])
    for i in range(n + 1):
        ops.append(f"{x0} {y_top - i * rh} m {x0 + m * cw} {y_top - i * rh} l S")
    for j in range(m + 1):
        ops.append(f"{x0 + j * cw} {y_top} m {x0 + j * cw} {y_top - n * rh} l S")
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            ops.append(f"BT /F1 10 Tf {x0 + j * cw + 4} {y_top - (i + 1) * rh + 6} Td "
                       f"({cell}) Tj ET")
    stream = "\n".join(ops).encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objs) + 1, xref))
    return out.getvalue()


PDF_ROWS = [["Region", "Revenue"], ["North", "120"], ["South", "90"], ["East", "60"]]


# --------------------------------------------------------------------------
# Parsing -- no DB
# --------------------------------------------------------------------------


def _profiled(filename, data, text=""):
    raw = uploads.raw_tables(filename, data, text, max_pdf_pages=5)
    return [t for t in (table_store_profile(r) for r in raw) if t is not None]


def table_store_profile(raw):
    from app.doctables import extract

    return extract.profile(raw)


def test_a_csv_upload_is_one_table():
    [table] = _profiled("sales.csv", SALES_CSV.encode())
    assert [c.type for c in table.columns] == ["category", "number"]


def test_every_excel_sheet_is_a_table_named_after_it():
    tables = _profiled("book.xlsx", _xlsx())
    assert [t.name for t in tables] == ["book.xlsx · Sales", "book.xlsx · Costs"]
    # 90.0 is read as 90, and the blank row is skipped rather than a row.
    assert len(tables[0].cells) == 3


def test_excel_text_is_what_the_model_is_shown():
    text, truncated = extract_text("book.xlsx", _xlsx(), max_chars=10_000,
                                   max_pdf_pages=5, max_csv_rows=100)
    assert "## Sales" in text and "South | 90" in text and not truncated


def test_word_tables_in_an_upload():
    [table] = _profiled("budget.docx", _docx())
    assert [c["name"] for c in (c.as_dict() for c in table.columns)] == ["Team", "Budget"]


def test_a_table_a_pdf_draws_is_read_with_its_columns():
    [table] = _profiled("report.pdf", ruled_pdf(PDF_ROWS))
    assert table.name == "report.pdf (page 1)"
    assert [c.type for c in table.columns] == ["category", "number"]


def test_pipe_tables_in_a_markdown_upload():
    md = "| Team | Budget |\n|---|---|\n| Core | 10 |\n| Growth | 5 |\n"
    [table] = _profiled("plan.md", md.encode(), md)
    assert len(table.cells) == 2


def test_a_file_with_no_table_has_none():
    assert _profiled("notes.txt", b"just words", "just words") == []


def test_an_upload_is_offered_as_the_askers_file():
    ref = table_store.TableRef(
        id="t", name="book.xlsx · Sales", columns=({"key": "c0", "name": "Region",
                                                     "type": "category"},),
        row_count=3, truncated=False, notes=(), document_title="book.xlsx",
        provider="upload", source_uri=None, attachment_id="a",
    )
    line = doc_tables.catalogue([ref])
    assert "the asker uploaded to this chat" in line and "[upload]" not in line


def test_starters_come_from_the_files_own_headers():
    from app.api.chat import upload_starters

    ref = table_store.TableRef(
        id="t", name="s", row_count=3, truncated=False, notes=(),
        columns=({"key": "c0", "name": "Region", "type": "category"},
                 {"key": "c1", "name": "Revenue", "type": "number"}),
        document_title="book.xlsx", provider="upload", source_uri=None, attachment_id="a",
    )
    assert upload_starters([ref]) == ["Total Revenue by Region"]


# --------------------------------------------------------------------------
# Storage, privacy and Chart mode -- a real database
# --------------------------------------------------------------------------


@pytest.fixture
def blob(monkeypatch):
    fake = FakeBlobStore()
    monkeypatch.setattr(attach_store, "blobstore", fake)
    return fake


def _user(conn, org):
    return str(conn.execute(
        "INSERT INTO users (email, org_id, role) VALUES (%s, %s, 'member') RETURNING id",
        (f"{uuid.uuid4().hex[:8]}@acme.test", org),
    ).fetchone()[0])


def _chat(conn, org, user):
    return str(conn.execute(
        "INSERT INTO conversations (org_id, user_id) VALUES (%s, %s) RETURNING id",
        (org, user),
    ).fetchone()[0])


@pytest.fixture
def chat(org_cleanup):
    with get_connection() as conn:
        org = str(conn.execute(
            "INSERT INTO organizations (name) VALUES (%s) RETURNING id",
            (f"uploads-{uuid.uuid4().hex[:8]}",),
        ).fetchone()[0])
        user, other = _user(conn, org), _user(conn, org)
        conv, conv2 = _chat(conn, org, user), _chat(conn, org, user)
        conn.commit()
    org_cleanup.append(org)
    return {"org": org, "user": user, "other": other, "conv": conv, "conv2": conv2}


def _upload(chat, blob, filename, data, text=None):
    text = text if text is not None else extract_text(
        filename, data, max_chars=10_000, max_pdf_pages=5, max_csv_rows=100)[0]
    saved = attach_store.save_attachment(
        org_id=chat["org"], conversation_id=chat["conv"], user_id=chat["user"],
        filename=filename, content_type="application/octet-stream", content=text,
        truncated=False, data=data,
    )
    uploads.read_upload_tables(saved.id, org_id=chat["org"], user_id=chat["user"],
                               filename=filename, data=data, text=text, max_pdf_pages=5)
    return saved


SANA = Viewer(email="sana@acme.test")


@requires_db
def test_an_upload_is_private_to_its_uploader_and_chat(chat, blob):
    _upload(chat, blob, "sales.csv", SALES_CSV.encode())
    mine = UploadScope(conversation_id=chat["conv"], user_id=chat["user"])
    [table] = table_store.list_upload_tables(org_id=chat["org"], uploads=mine)
    assert table.is_upload and table.document_title == "sales.csv"

    # Another member, the same person in another chat, and the document
    # listing: none of them see it, and a run-time lookup is "not found".
    theirs = UploadScope(conversation_id=chat["conv"], user_id=chat["other"])
    elsewhere = UploadScope(conversation_id=chat["conv2"], user_id=chat["user"])
    for scope in (theirs, elsewhere):
        assert table_store.list_upload_tables(org_id=chat["org"], uploads=scope) == []
        assert table_store.get_table(table.id, org_id=chat["org"], workspace_id=None,
                                     viewer=SANA, uploads=scope) is None
    assert table_store.list_tables(org_id=chat["org"], workspace_id=None, viewer=SANA) == []
    assert table_store.get_table(table.id, org_id=chat["org"], workspace_id=None,
                                 viewer=SANA) is None
    assert table_store.get_table(table.id, org_id=chat["org"], workspace_id=None,
                                 viewer=SANA, uploads=mine) is not None


@requires_db
def test_tables_cannot_be_hung_on_someone_elses_file(chat, blob):
    saved = _upload(chat, blob, "notes.txt", b"no table here", "no table here")
    from app.doctables import extract

    table = extract.profile(extract.parse_csv(SALES_CSV, name="x"))
    assert table_store.replace_attachment_tables(
        saved.id, org_id=chat["org"], user_id=chat["other"], tables=[table]) == 0


@requires_db
def test_removing_the_file_removes_its_tables(chat, blob):
    saved = _upload(chat, blob, "book.xlsx", _xlsx())
    mine = UploadScope(conversation_id=chat["conv"], user_id=chat["user"])
    assert len(table_store.list_upload_tables(org_id=chat["org"], uploads=mine)) == 2
    attach_store.delete_attachment(attachment_id=saved.id, org_id=chat["org"],
                                   conversation_id=chat["conv"], user_id=chat["user"])
    with get_connection() as conn:
        left = conn.execute("SELECT count(*) FROM doc_tables WHERE attachment_id = %s",
                            (saved.id,)).fetchone()[0]
    assert left == 0


@requires_db
def test_an_uploaded_sheet_is_summed_by_its_column(chat, blob):
    _upload(chat, blob, "report.pdf", ruled_pdf(PDF_ROWS))
    mine = UploadScope(conversation_id=chat["conv"], user_id=chat["user"])
    [table] = table_store.list_upload_tables(org_id=chat["org"], uploads=mine)
    spec = ChartSpec(metric=TABLE_METRIC, group_by="c0", period="month", chart="bar",
                     table_id=table.id, measure="sum", value="c1")
    response = insights_agent.InsightsAgent().answer(
        "revenue by region", chat["org"], conversation_id=chat["conv"],
        user_id=chat["user"], viewer=SANA, spec=spec,
    )
    assert {p["group"]: p["value"] for p in response.chart["points"]} == {
        "North": 120.0, "South": 90.0, "East": 60.0}
    assert "the file you uploaded" in response.chart["caveat"]
    # The same spec from another member's chat is not a chart.
    other = insights_agent.InsightsAgent().answer(
        "revenue by region", chat["org"], conversation_id=chat["conv"],
        user_id=chat["other"], viewer=SANA, spec=spec,
    )
    assert other.chart is None


class _FigureReader:
    """The text adapter's surface, counting calls."""

    origin = "text"

    def __init__(self):
        self.calls = 0

    def wants(self, doc):
        return True

    def extract(self, doc):
        from app.doctables import extract

        self.calls += 1
        return [extract.profile(extract.parse_csv(SALES_CSV, name=doc.title))]


@requires_db
def test_a_file_without_a_table_is_read_for_figures_once(chat, blob):
    _upload(chat, blob, "memo.txt", b"North made 120, South 90.", "North made 120, South 90.")
    mine = UploadScope(conversation_id=chat["conv"], user_id=chat["user"])
    reader = _FigureReader()
    assert uploads.read_upload_figures(org_id=chat["org"], uploads=mine, adapter=reader) == 1
    assert uploads.read_upload_figures(org_id=chat["org"], uploads=mine, adapter=reader) == 0
    assert reader.calls == 1
    [table] = table_store.list_upload_tables(org_id=chat["org"], uploads=mine)
    assert table.origin == "text"


@requires_db
def test_chart_mode_offers_the_chats_files_first(chat, blob, monkeypatch):
    _upload(chat, blob, "sales.csv", SALES_CSV.encode())
    monkeypatch.setattr(routing, "_document_tables", lambda *a, **k: ["doc-table"])
    monkeypatch.setattr(uploads, "read_upload_figures", lambda **k: 0)
    mine = UploadScope(conversation_id=chat["conv"], user_id=chat["user"])
    offered = routing._chartable_tables(chat["org"], None, SANA, "revenue", mine)
    assert offered[0].document_title == "sales.csv" and offered[1:] == ["doc-table"]
    # No viewer still means no tables at all, uploads included.
    assert routing._chartable_tables(chat["org"], None, None, "revenue", mine) == []


def test_a_line_on_a_table_runs_over_its_date_column():
    """"Revenue per month for North as a line" drew one bar: the model left
    group_by empty, so the period had no axis to ride on."""
    from app.doctables import extract

    raw = extract.parse_csv("Date,Region,Revenue\n2026-01-05,North,10\n2026-02-07,North,20\n"
                            "2026-02-09,South,5\n", name="s")
    profiled = extract.profile(raw)
    ref = table_store.TableRef(
        id="t", name="s", columns=tuple(c.as_dict() for c in profiled.columns),
        row_count=3, truncated=False, notes=(), document_title="s.csv",
        provider="attachment", source_uri=None, attachment_id="a",
    )
    pick = doc_tables.parse_pick(
        {"table": "T1", "measure": "sum", "value": "Revenue", "chart": "line",
         "filters": {"Region": "North"}}, {"T1": ref})
    assert pick.group_by == "c0" and pick.chart == "line"
