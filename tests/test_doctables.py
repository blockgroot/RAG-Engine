"""Tables inside documents (`app/doctables`, `insights/tables.py`).

Step 3 of docs/plans/2026-09-30-open-ended-charts.md: "chart last year's
revenue by region" from a sales sheet. Pinned here:

- parsing is deterministic and honest about what it could not read (Indian
  digit grouping, currency, day-first dates, unparseable cells left OUT of a
  sum, never counted as zero);
- a spreadsheet embeds a DESCRIPTION of its columns, never its figures;
- a table is exactly as visible as its document, checked when it is offered
  AND again when it is run, and no viewer means no table chart at all;
- the model picks a table by handle and columns by header text, and a pick
  that does not fit the table is refused with the options.
"""

from __future__ import annotations

import json
import uuid
from io import BytesIO

import pytest

from app.agent import insights_agent
from app.db import get_connection
from app.doctables import extract, store as table_store
from app.ingestion.pipeline import _store_tables
from app.insights import resolve
from app.insights import tables as doc_tables
from app.insights.resolve import CannotChart, ChartSpec, TABLE_METRIC
from app.sources.base import SourceDocument
from app.sources.google_drive import _extract_docx_text, _table_document
from app.vectorstore.base import Viewer
from .conftest import requires_db

SALES_CSV = (
    "Month,Region,Revenue,Notes\n"
    "2025-01-15,North,\"₹1,20,000\",good\n"
    "2025-01-20,South,\"₹90,000\",\n"
    "2025-02-10,North,\"₹1,40,000\",\n"
    "2025-02-12,South,N/A,missing\n"
)


# --------------------------------------------------------------------------
# Parsing -- no DB
# --------------------------------------------------------------------------


@pytest.mark.parametrize("text,number,unit", [
    ("₹1,20,000", 120000.0, "₹"), ("(1,200)", -1200.0, ""), ("12.5%", 12.5, "%"),
    ("$3.2k", 3200.0, "$"), ("2 lakh", 200000.0, ""), ("Rs. 500", 500.0, "₹"),
    ("1.5 cr", 15000000.0, ""), ("N/A", None, ""), ("", None, ""),
])
def test_numbers_as_spreadsheets_write_them(text, number, unit):
    assert extract.parse_number(text) == (number, unit)


def test_a_column_is_typed_as_a_whole():
    table = extract.profile(extract.parse_csv(SALES_CSV, name="Sales"))
    kinds = {c.name: (c.type, c.unit) for c in table.columns}
    assert kinds == {"Month": ("date", ""), "Region": ("category", ""),
                     "Revenue": ("number", "₹"), "Notes": ("category", "")}
    revenue = next(c for c in table.columns if c.name == "Revenue")
    assert revenue.unparsed == 1  # "N/A" -- stored raw, left out of any sum
    assert "c2" not in table.cells[3]


def test_dates_are_read_day_first_when_the_column_allows_both():
    raw = extract.RawTable("t", ["When", "N"], [["03/04/2025", "1"], ["13/04/2025", "2"]])
    table = extract.profile(raw)
    assert table.cells[0]["c0"] == "2025-04-03"


def test_a_year_column_is_a_date_not_a_number_to_add_up():
    raw = extract.RawTable("t", ["Year", "Sales"], [["2023", "10"], ["2024", "12"]])
    assert [c.type for c in extract.profile(raw).columns] == ["date", "number"]


def test_markdown_tables_are_named_after_their_heading():
    text = ("Intro\n\n## Q1 budget\n\n| Team | Budget |\n|---|---|\n| Core | $10k |\n"
            "| Growth | $5k |\n\nNot | a table\n")
    [found] = extract.find_markdown_tables(text, title="Plan")
    assert found.name == "Q1 budget"
    assert found.rows == [["Core", "$10k"], ["Growth", "$5k"]]


def test_a_one_row_table_is_not_a_table():
    assert extract.find_markdown_tables("| a | b |\n|---|---|\n| 1 | 2 |\n", title="x") == []


def test_parsing_is_bounded_and_says_so(monkeypatch):
    monkeypatch.setattr(extract, "MAX_ROWS", 3)
    rows = "\n".join(f"r{i},{i}" for i in range(10))
    raw = extract.parse_csv("name,n\n" + rows, name="big")
    assert raw.truncated and len(raw.rows) <= 4
    assert any("first 3 rows" in n for n in extract.profile(raw).notes)


def test_a_spreadsheet_embeds_its_columns_never_its_figures():
    content, tables = _table_document(SALES_CSV, name="Sales 2025", sheet=True)
    assert "Revenue" in content and "North" in content
    for figure in ("1,20,000", "120000", "90,000", "1,40,000"):
        assert figure not in content
    assert tables and "first tab" in tables[0].notes[0]


def test_word_tables_are_no_longer_dropped():
    import docx

    document = docx.Document()
    document.add_paragraph("Budget below.")
    grid = document.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Team", "Budget"), ("Core", "10"), ("Growth", "5")]):
        grid.rows[r].cells[0].text, grid.rows[r].cells[1].text = a, b
    data = BytesIO()
    document.save(data)
    text = _extract_docx_text(data.getvalue())
    assert "Budget below." in text
    [found] = extract.find_markdown_tables(text, title="Doc")
    assert found.header == ["Team", "Budget"]


# --------------------------------------------------------------------------
# Offering and validating -- no DB
# --------------------------------------------------------------------------


def _ref(**kw):
    base = dict(
        id=str(uuid.uuid4()), name="Sales 2025", row_count=4, truncated=False,
        notes=(), document_title="Sales 2025", provider="google", source_uri=None,
        columns=(
            {"key": "c0", "name": "Month", "type": "date"},
            {"key": "c1", "name": "Region", "type": "category", "samples": ["North", "South"]},
            {"key": "c2", "name": "Revenue", "type": "number", "unit": "₹"},
            {"key": "c3", "name": "Notes", "type": "text"},
        ),
    )
    base.update(kw)
    return table_store.TableRef(**base)


class FakeLLM:
    model = "t"
    last_usage = None

    def __init__(self, reply):
        self.reply, self.prompts = reply, []

    def generate(self, prompt, *, max_tokens=None):
        self.prompts.append(prompt)
        return json.dumps(self.reply)


def test_tables_are_offered_by_meaning_closest_first():
    """No stop-word list: each table's DOCUMENT is compared with the question
    in the retrieval index. Below the gate is ORDER, never removal: a Notion
    page that is mostly a table scored under it and "chart the quarterly
    budget page" was offered no table at all."""
    sales = _ref(document_id="d-sales")
    hiring = _ref(name="Hiring plan", document_title="Hiring plan", document_id="d-hiring",
                  columns=({"key": "c0", "name": "Role", "type": "category"},))
    scores = {"d-sales": 0.71, "d-hiring": 0.22}
    assert doc_tables.rank([hiring, sales], scores, floor=0.35) == [sales, hiring]
    assert doc_tables.rank([hiring, sales], {"d-sales": 0.1}, floor=0.35) == [sales, hiring]


def test_unmeasurable_similarity_offers_the_newest_tables():
    tables = [_ref(name=f"T{i}", document_id=f"d{i}") for i in range(9)]
    assert doc_tables.rank(tables, None, floor=0.35) == tables[:doc_tables.MAX_OFFERED]


def test_a_table_pick_becomes_a_spec_with_column_keys():
    sales = _ref()
    llm = FakeLLM({"intent": "chart", "metric": None, "table": "T1",
                   "group_by": "region", "value": "Revenue", "measure": "sum"})
    intent = resolve.classify_question("chart revenue by region", providers=["google"],
                                       llm=llm, fail_open=False, tables=[sales])
    assert "T1: \"Sales 2025\"" in llm.prompts[0]
    spec = intent.spec
    assert (spec.metric, spec.table_id, spec.group_by, spec.value, spec.measure) == (
        TABLE_METRIC, sales.id, "c1", "c2", "sum")
    assert spec.chart == "bar"


def test_a_date_breakdown_is_a_line():
    llm = FakeLLM({"intent": "chart", "table": "T1", "group_by": "Month", "value": "Revenue"})
    intent = resolve.classify_question("revenue per month chart", providers=["google"],
                                       llm=llm, fail_open=False, tables=[_ref()])
    assert intent.spec.chart == "line" and intent.spec.measure == "sum"


@pytest.mark.parametrize("reply,needle", [
    ({"table": "T1", "value": "Region", "measure": "sum"}, "can't be the value"),
    ({"table": "T1", "group_by": "Notes"}, "can't be the breakdown"),
    ({"table": "T1", "group_by": "Profit"}, "no column called"),
    ({"table": "T1", "measure": "sum"}, "number column"),
])
def test_a_pick_that_does_not_fit_is_refused_with_the_options(reply, needle):
    llm = FakeLLM({"intent": "chart", **reply})
    intent = resolve.classify_question("chart revenue", providers=["google"], llm=llm,
                                       fail_open=False, tables=[_ref()])
    assert intent.kind == "refuse" and needle in intent.message


def test_an_invented_handle_is_not_a_table():
    llm = FakeLLM({"intent": "chart", "metric": None, "table": "T9"})
    intent = resolve.classify_question("chart revenue", providers=["google"], llm=llm,
                                       fail_open=False, tables=[_ref()])
    assert intent.kind == "refuse" and intent.spec is None


def test_no_tables_means_no_table_section_in_the_prompt():
    llm = FakeLLM({"intent": "qa"})
    resolve.classify_question("chart revenue", providers=["google"], llm=llm,
                              fail_open=False)
    assert "TABLES INSIDE DOCUMENTS" not in llm.prompts[0]


def test_a_table_spec_survives_the_graph_state():
    spec = ChartSpec(metric=TABLE_METRIC, group_by="c1", period="month", chart="bar",
                     measure="sum", table_id="abc", value="c2", filters=(("c1", "North"),))
    assert resolve.spec_from_dict(json.loads(json.dumps(resolve.spec_to_dict(spec)))) == spec


def test_a_table_chart_without_a_viewer_is_refused():
    spec = ChartSpec(metric=TABLE_METRIC, group_by=None, period="month", chart="bar",
                     table_id=str(uuid.uuid4()))
    with pytest.raises(CannotChart):
        insights_agent.run_table_spec(spec, org_id="o", workspace_id=None, viewer=None)


# --------------------------------------------------------------------------
# Storage, access and SQL -- a real database
# --------------------------------------------------------------------------


@pytest.fixture
def org(org_cleanup):
    with get_connection() as conn:
        row = conn.execute(
            "INSERT INTO organizations (name) VALUES (%s) RETURNING id",
            (f"tables-{uuid.uuid4().hex[:8]}",),
        ).fetchone()
        conn.commit()
    org_cleanup.append(str(row[0]))
    return str(row[0])


def _document(org_id, *, public=True, viewers=None, title="Sales 2025"):
    with get_connection() as conn:
        row = conn.execute(
            """
            INSERT INTO documents (org_id, title, source_uri, source_provider,
                                   source_external_id, doc_is_public, doc_viewers)
            VALUES (%s, %s, 'https://drive.test/x', 'google', %s, %s, %s)
            RETURNING id
            """,
            (org_id, title, uuid.uuid4().hex, public, viewers or []),
        ).fetchone()
        conn.commit()
    return str(row[0])


def _ingest(org_id, document_id, text=SALES_CSV):
    raw = extract.parse_csv(text, name="Sales 2025")
    doc = SourceDocument(external_id="x", title="Sales 2025", content="", tables=[raw])
    _store_tables(document_id, doc, org_id=org_id, workspace_id=None)


SANA = Viewer(email="sana@acme.test")


@requires_db
def test_ingest_stores_typed_rows(org):
    doc = _document(org)
    _ingest(org, doc)
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    assert (table.name, table.row_count) == ("Sales 2025", 4)
    assert [c["type"] for c in table.columns] == ["date", "category", "number", "category"]


@requires_db
def test_markdown_tables_are_captured_from_content(org):
    doc = _document(org, title="Plan")
    md = "## Budget\n\n| Team | Budget |\n|---|---|\n| Core | 10 |\n| Growth | 5 |\n"
    _store_tables(doc, SourceDocument(external_id="p", title="Plan", content=md),
                  org_id=org, workspace_id=None)
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    assert table.name == "Budget"


@requires_db
def test_a_table_is_exactly_as_visible_as_its_document(org):
    doc = _document(org, public=False, viewers=["finance@acme.test"])
    _ingest(org, doc)
    assert table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA) == []
    assert table_store.list_tables(org_id=org, workspace_id=None,
                                   viewer=Viewer.public_only_viewer()) == []
    finance = Viewer(email="finance@acme.test")
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=finance)
    # ...and again at RUN time: a spec for it, run as Sana, is refused.
    spec = ChartSpec(metric=TABLE_METRIC, group_by="c1", period="month", chart="bar",
                     measure="sum", value="c2", table_id=table.id)
    with pytest.raises(CannotChart, match="no longer available"):
        insights_agent.run_table_spec(spec, org_id=org, workspace_id=None, viewer=SANA)


@requires_db
def test_a_table_is_pinned_to_its_scope(org):
    doc = _document(org)
    _ingest(org, doc)
    with get_connection() as conn:
        user = conn.execute(
            "INSERT INTO users (email, org_id, role) VALUES (%s, %s, 'member') RETURNING id",
            (f"{uuid.uuid4().hex[:8]}@acme.test", org),
        ).fetchone()
        space = conn.execute(
            "INSERT INTO workspaces (org_id, name, created_by) VALUES (%s, 'S', %s) RETURNING id",
            (org, user[0]),
        ).fetchone()
        conn.commit()
    assert table_store.list_tables(org_id=org, workspace_id=str(space[0]), viewer=SANA) == []


@requires_db
def test_revenue_by_region_is_summed_from_the_cells(org):
    doc = _document(org)
    _ingest(org, doc)
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    spec = ChartSpec(metric=TABLE_METRIC, group_by="c1", period="month", chart="bar",
                     measure="sum", value="c2", table_id=table.id)
    panel, _ = insights_agent.run_table_spec(spec, org_id=org, workspace_id=None, viewer=SANA)
    assert {p["group"]: p["value"] for p in panel["points"]} == {"North": 260000, "South": 90000}
    assert panel["title"] == "Sum of Revenue by Region — Sales 2025"
    assert panel["unit"] == "₹"
    assert "1 Revenue cells were not numbers" in panel["caveat"]
    assert {d["attrs"]["c1"] for d in panel["details"]} == {"North", "South"}


@requires_db
def test_a_date_column_is_the_time_axis(org):
    doc = _document(org)
    _ingest(org, doc)
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    spec = ChartSpec(metric=TABLE_METRIC, group_by="c0", period="month", chart="line",
                     measure="sum", value="c2", table_id=table.id,
                     filters=(("c1", "north"),))
    panel, _ = insights_agent.run_table_spec(spec, org_id=org, workspace_id=None, viewer=SANA)
    assert [(p["bucket"][:7], p["value"]) for p in panel["points"]] == [
        ("2025-01", 120000), ("2025-02", 140000)]
    assert "Region: North" in panel["title"]


@requires_db
def test_an_unknown_filter_value_is_refused_listing_what_exists(org):
    doc = _document(org)
    _ingest(org, doc)
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    spec = ChartSpec(metric=TABLE_METRIC, group_by=None, period="month", chart="bar",
                     table_id=table.id, filters=(("c1", "Mars"),))
    with pytest.raises(CannotChart) as err:
        insights_agent.run_table_spec(spec, org_id=org, workspace_id=None, viewer=SANA)
    assert "North" in str(err.value) and "South" in str(err.value)


@requires_db
def test_reingest_replaces_rather_than_appends(org):
    doc = _document(org)
    _ingest(org, doc)
    _ingest(org, doc, text="Region,Revenue\nEast,5\nWest,7\n")
    [table] = table_store.list_tables(org_id=org, workspace_id=None, viewer=SANA)
    assert table.row_count == 2
