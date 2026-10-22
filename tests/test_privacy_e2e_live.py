"""End to end: no Aadhaar, PAN, card, mobile, IFSC or UPI value leaves the process.

LIVE: the real chat API and Slack answer path, the real model (LLM_* from the
environment), the real embedder and reranker, the Groq checks when switched on,
and the staging database. Every HTTP request leaving the process is recorded
and searched for the raw values; answers must still carry the real values.

Opt-in only (costs model quota and several minutes):

    PRIVACY_E2E=1 GUARD_MODE=shadow GUARD_ANSWER_CHECK=true \\
    INGEST_CONTEXTUAL_ENABLED=true INGEST_CONTEXTUAL_DEFER=false \\
    PERSONAL_MEMORY_ENABLED=true WEB_SEARCH_ENABLED=true \\
    pytest tests/test_privacy_e2e_live.py -s

Scenarios: uploads (.txt .csv .xlsx .docx .pdf .md), a follow-up that reads
chat history, Deep mode, Chart mode, an indexed document (ingest,
contextualization, guard scoring, embedding, reranking), a question that
contains an identifier, a look-alike order id that must NOT be masked,
personal memory, the web-search decision, the Slack answer path and an Agent
mode draft (tool call). Web search sends the asker's own words by design; what
it was sent is reported, not asserted.
"""

from __future__ import annotations

import io
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace

import httpx
import pytest
from stdnum import verhoeff

from app.attachments import store as attach_store
from app.db import get_connection

from .conftest import requires_db
from .test_attachment_storage import FakeBlobStore

pytestmark = [
    requires_db,
    pytest.mark.skipif(os.getenv("PRIVACY_E2E") != "1", reason="live: set PRIVACY_E2E=1"),
]

PAUSE = float(os.getenv("E2E_PAUSE", "12"))  # a free model tier allows ~15 requests a minute


def _aadhaar(body: str) -> str:
    return body + verhoeff.calc_check_digit(body)


# Synthetic values: each passes its real check digit / validity rule.
RAVI = {
    "mobile": "+91 98765 43210",
    "pan": "AJKPK4821Q",
    "ifsc": "HDFC0004521",
    "upi": "ravi.kumar@okhdfc",
    "aadhaar": _aadhaar("49183726501"),
}
PRIYA = {"mobile": "9123456780", "pan": "BQRPS7310L", "upi": "priya.s@ybl"}
MEENA = {"aadhaar": _aadhaar("73829104562"), "vid": _aadhaar("918273645501928"), "pan": "CMNPM5512D"}
ARJUN = {"card": "4111 1111 1111 1111", "mobile": "+91-99887-76655"}
INDEXED = {"mobile": "9988012345", "pan": "DXYPA9087K", "ifsc": "ICIC0001122"}
ASKER = {"mobile": "9876501234"}
ORDER_ID = "5012345678"  # 10 digits, starts with 5: an order id, never masked

ALL_VALUES = [*RAVI.values(), *PRIYA.values(), *MEENA.values(), *ARJUN.values(),
              *INDEXED.values(), *ASKER.values()]
SAFE_HOSTS = {"testserver"}  # our own app, called by the test client
WEB = "duckduckgo"          # the asker's own words, by design: reported, not asserted


def _spellings(value: str) -> set[str]:
    digits = re.sub(r"\D", "", value)
    out = {value, value.lower(), value.upper()}
    if len(digits) >= 10:
        out.add(digits)
        out.add(digits[-10:])
    return {s for s in out if len(s) >= 8}


LEAK_PATTERNS = {v: _spellings(v) for v in ALL_VALUES}
REPORT: list[tuple[str, str, str]] = []


# -- fixtures -----------------------------------------------------------------------


@pytest.fixture(scope="module")
def outbound():
    """Every request leaving the process: sync and async httpx, and web search."""
    sent: list[tuple[str, str]] = []
    real_send, real_async = httpx.Client.send, httpx.AsyncClient.send

    def _body(request) -> str:
        try:
            return request.read().decode("utf-8", "replace")
        except Exception:  # noqa: BLE001 - a streaming body we cannot read
            return ""

    def spy(self, request, **kw):
        sent.append((request.url.host, _body(request)))
        return real_send(self, request, **kw)

    async def aspy(self, request, **kw):
        sent.append((request.url.host, _body(request)))
        return await real_async(self, request, **kw)

    from app.websearch import duckduckgo

    real_search = duckduckgo.DuckDuckGoSearch.search

    def search_spy(self, query, *a, **k):
        sent.append((WEB, query))
        return real_search(self, query, *a, **k)

    mp = pytest.MonkeyPatch()
    mp.setattr(httpx.Client, "send", spy)
    mp.setattr(httpx.AsyncClient, "send", aspy)
    mp.setattr(duckduckgo.DuckDuckGoSearch, "search", search_spy)
    yield sent
    mp.undo()


@pytest.fixture(scope="module")
def ctx(outbound):
    from fastapi.testclient import TestClient

    from app.api.main import create_app
    from app.auth import create_admin, create_session_token

    mp = pytest.MonkeyPatch()
    mp.setenv("AUTH_JWT_SECRET", "test-jwt-secret-do-not-use-in-prod")
    mp.setenv("CHAT_STREAM_WORD_DELAY_MS", "0")
    mp.setattr(attach_store, "blobstore", FakeBlobStore())  # no uploads to Cloudinary

    with get_connection() as conn:
        org = str(conn.execute(
            "INSERT INTO organizations (name) VALUES (%s) RETURNING id",
            (f"privacy-e2e-{uuid.uuid4().hex[:8]}",),
        ).fetchone()[0])
        conn.commit()
    email = f"sana.asiwal+privacy-e2e-{uuid.uuid4().hex[:6]}@syvora.com"
    user = create_admin(email, org)
    try:
        yield SimpleNamespace(
            client=TestClient(create_app()), cookies={"session": create_session_token(user)},
            org=org, user=user, email=email,
        )
    finally:
        with get_connection() as conn:
            conn.execute("DELETE FROM organizations WHERE id = %s::uuid", (org,))
            conn.execute("DELETE FROM users WHERE email = %s", (email,))
            conn.commit()
        mp.undo()


@pytest.fixture(autouse=True)
def _no_leaks_after_each(outbound):
    start = len(outbound)
    yield
    leaks = _leaks(outbound[start:])
    assert not leaks, f"raw identifiers left the process: {leaks}"


# -- files --------------------------------------------------------------------------


def _txt() -> bytes:
    return (
        "Vendor onboarding record\n"
        "Vendor contact: Ravi Kumar\n"
        f"Mobile: {RAVI['mobile']}\n"
        f"PAN: {RAVI['pan']}\n"
        f"Bank IFSC: {RAVI['ifsc']}\n"
        f"UPI for payments: {RAVI['upi']}\n"
        f"Aadhaar (KYC): {RAVI['aadhaar'][:4]} {RAVI['aadhaar'][4:8]} {RAVI['aadhaar'][8:]}\n"
        f"Laptop order number: {ORDER_ID}\n"
        "Contact email: ravi.kumar@example.com\n"
    ).encode()


def _csv() -> bytes:
    return (
        "name,team,mobile,pan,upi\n"
        f"Priya Shah,Design,{PRIYA['mobile']},{PRIYA['pan']},{PRIYA['upi']}\n"
        "Karan Mehta,Sales,,,\n"
        "Leela Das,Design,,,\n"
    ).encode()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee", "Team", "Aadhaar", "Virtual ID", "PAN"])
    ws.append(["Meena Iyer", "Finance", MEENA["aadhaar"], MEENA["vid"], MEENA["pan"]])
    ws.append(["Rohit Jain", "Finance", "", "", ""])
    ws.append(["Asha Rao", "Ops", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Travel reimbursement for Arjun Rao.")
    doc.add_paragraph(f"Corporate card used: {ARJUN['card']}. Reach Arjun on {ARJUN['mobile']}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    fig = plt.figure(figsize=(8.5, 3))
    fig.text(0.05, 0.7, f"KYC summary: Meena Iyer, PAN {MEENA['pan']}", fontsize=12)
    fig.text(0.05, 0.5, f"Aadhaar {MEENA['aadhaar']}", fontsize=12)
    buf = io.BytesIO()
    fig.savefig(buf, format="pdf")
    plt.close(fig)
    return buf.getvalue()


def _md() -> bytes:
    return f"# Support escalation\n\nCall the on-call lead Arjun at {ARJUN['mobile']}.\n".encode()


# -- helpers ------------------------------------------------------------------------


def _conversation(ctx) -> str:
    with get_connection() as conn:
        conv = str(conn.execute(
            "INSERT INTO conversations (org_id, user_id) VALUES (%s, %s) RETURNING id",
            (ctx.org, ctx.user.id),
        ).fetchone()[0])
        conn.commit()
    return conv


def _upload(ctx, conv: str, files: list[tuple[str, bytes]]) -> dict:
    r = ctx.client.post(
        f"/chat/conversations/{conv}/attachments",
        files=[("file", (name, body, "application/octet-stream")) for name, body in files],
        cookies=ctx.cookies,
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert len(out["attachments"]) == len(files) and not out["rejected"], out
    return out


def _stream(ctx, body: dict) -> dict:
    r = ctx.client.post("/chat/stream", json=body, cookies=ctx.cookies)
    assert r.status_code == 200, r.text
    done: dict = {}
    for block in r.text.strip().split("\n\n"):
        lines = block.split("\n")
        if lines and lines[0] in ("event: done", "event: error"):
            data = json.loads("\n".join(l[6:] for l in lines[1:] if l.startswith("data: ")))
            done = data if lines[0] == "event: done" else {"answer": "", "error": data.get("message")}
    return done


def _ask(ctx, question: str, conv: str | None = None, mode: str | None = None) -> dict:
    body: dict = {"question": question}
    if conv:
        body["conversation_id"] = conv
    if mode:
        body["mode"] = mode
    done = _stream(ctx, body)
    if done.get("error"):  # one retry: a busy model (503) is not what this tests
        time.sleep(PAUSE * 2)
        done = _stream(ctx, body)
    time.sleep(PAUSE)
    return done


def _leaks(sent: list[tuple[str, str]]) -> list[tuple[str, str]]:
    found = []
    for host, body in sent:
        if host in SAFE_HOSTS or host == WEB:
            continue
        for value, spellings in LEAK_PATTERNS.items():
            if any(s in body for s in spellings):
                found.append((host, value))
    return found


def _has(answer: str, value: str) -> bool:
    digits = re.sub(r"\D", "", value)
    if len(digits) >= 10:
        return digits[-10:] in re.sub(r"\D", "", answer)
    return value.lower() in answer.lower()


def _check(label: str, done_or_text, expected: list[str]) -> None:
    if isinstance(done_or_text, dict):
        answer = done_or_text.get("answer") or ""
        err = done_or_text.get("error")
    else:
        answer, err = done_or_text or "", None
    missing = [e for e in expected if not _has(answer, e)]
    REPORT.append((label, "OK" if not missing else "MISSING " + ", ".join(missing),
                   answer[:240] or f"ERROR: {err}"))


# -- scenarios ----------------------------------------------------------------------


def test_uploads_questions_and_follow_up(ctx, outbound):
    conv = _conversation(ctx)
    _upload(ctx, conv, [
        ("vendor.txt", _txt()), ("team.csv", _csv()), ("kyc.xlsx", _xlsx()),
        ("travel.docx", _docx()), ("kyc.pdf", _pdf()),
    ])
    REPORT.append(("upload", "5 files saved (.txt .csv .xlsx .docx .pdf)", ""))
    cases = [
        ("What is Ravi Kumar's mobile number?", [RAVI["mobile"]]),
        ("And what is his PAN?", [RAVI["pan"]]),  # follow-up: reads chat history
        ("What is Ravi Kumar's bank IFSC code?", [RAVI["ifsc"]]),
        (f"Whose UPI ID is {RAVI['upi']}?", ["Ravi"]),
        ("What is the laptop order number for Ravi?", [ORDER_ID]),
        ("What is Priya Shah's mobile number and UPI ID in the team sheet?",
         [PRIYA["mobile"], PRIYA["upi"]]),
        ("What is Meena Iyer's PAN?", [MEENA["pan"]]),
        ("How can I reach Arjun Rao?", [ARJUN["mobile"]]),
        ("What is Ravi Kumar's contact email?", ["ravi.kumar@example.com"]),
    ]
    for question, expected in cases:
        _check(question, _ask(ctx, question, conv), expected)
    assert any("<PHONE_" in b for _, b in outbound), "masking never engaged"


def test_deep_mode(ctx):
    conv = _conversation(ctx)
    _upload(ctx, conv, [("vendor.txt", _txt())])
    _check("[deep] Ravi's PAN and IFSC",
           _ask(ctx, "Give me Ravi Kumar's PAN and IFSC code.", conv, mode="deep"),
           [RAVI["pan"], RAVI["ifsc"]])


def test_chart_mode_on_a_sheet_with_identifiers(ctx):
    conv = _conversation(ctx)
    _upload(ctx, conv, [("kyc.xlsx", _xlsx())])
    done = _ask(ctx, "Draw a bar chart of the number of employees by team.", conv, mode="chart")
    REPORT.append(("[chart] employees by team", "chart drawn" if done.get("chart") else "no chart",
                   (done.get("answer") or "")[:200] or f"ERROR: {done.get('error')}"))


def test_markdown_upload(ctx):
    conv = _conversation(ctx)
    _upload(ctx, conv, [("oncall.md", _md())])
    _check("[.md] on-call number",
           _ask(ctx, "What number should I call to reach the on-call lead?", conv),
           [ARJUN["mobile"]])


def test_indexed_document(ctx, outbound):
    from app.ingestion.pipeline import ingest_source
    from app.sources.base import DocAccess, SourceDocument, SourceRef

    when = datetime(2026, 10, 1, tzinfo=timezone.utc)
    text = (
        "Facilities vendor directory.\n"
        f"The office caterer Anil Desai can be reached on {INDEXED['mobile']}. "
        f"Invoices go to PAN {INDEXED['pan']}, paid to IFSC {INDEXED['ifsc']}. "
        f"Slack him at <@U0456ANIL> for menu changes."
    )

    class _Adapter:
        def list_documents(self):
            return [SourceRef("vendors-1", "Facilities vendor directory", last_modified=when)]

        def fetch_document(self, external_id):
            return SourceDocument(
                external_id="vendors-1", title="Facilities vendor directory", content=text,
                source_uri="https://example.com/vendors-1", last_modified=when,
                access=DocAccess.scope_public(),
            )

        def get_last_modified(self, external_id):
            return when

    before = len(outbound)
    ingest_source(_Adapter(), ctx.org, provider="notion")
    hosts = sorted({h for h, _ in outbound[before:]})
    REPORT.append(("ingest", f"{len(outbound) - before} outbound calls", f"hosts {hosts}"))
    _check("[indexed] caterer phone",
           _ask(ctx, "What is the phone number of the office caterer Anil Desai?"),
           [INDEXED["mobile"]])
    _check("[indexed] PAN in the question",
           _ask(ctx, f"Who has the PAN {INDEXED['pan']}?"), ["Anil"])


def test_personal_memory(ctx):
    _check("[memory] asker's own number in the question",
           _ask(ctx, f"Remember that my mobile number is {ASKER['mobile']}. "
                     "What is the office caterer's IFSC code?"),
           [INDEXED["ifsc"]])


def test_web_search_decision(ctx, outbound):
    before = len(outbound)
    done = _ask(ctx, f"Search the web: which bank and branch uses the IFSC code {RAVI['ifsc']}?")
    web = [b for h, b in outbound[before:] if h == WEB]
    REPORT.append(("[web] IFSC typed by the asker",
                   f"web queries: {web or 'none'}",
                   (done.get("answer") or "")[:200] or f"ERROR: {done.get('error')}"))
    time.sleep(PAUSE)


def test_slack_answer_path(ctx):
    from app.api.slack_events import _answer

    reply = _answer("What is the phone number of the office caterer Anil Desai?", ctx.org, None, None)
    _check("[slack] caterer phone", getattr(reply, "text", "") or str(reply), [INDEXED["mobile"]])
    time.sleep(PAUSE)


def test_agent_mode_draft_tool_call(ctx):
    from app.actions import service
    from app.llm import build_llm_provider

    tool = {"type": "function", "function": {
        "name": "post_message",
        "description": "Post a message in a Slack channel.",
        "parameters": {"type": "object", "properties": {
            "channel": {"type": "string"},
            "message": {"type": "string"},
            "request_words": {"type": "string",
                              "description": "The asker's words that ask for this, exactly as written."},
        }, "required": ["channel", "message", "request_words"]},
    }}
    action = SimpleNamespace(tool=tool, lookup=False)
    question = (f"Post in #vendors that Ravi's new mobile is {RAVI['mobile']} "
                f"and his PAN is {RAVI['pan']}.")
    out = service._ask(build_llm_provider(), question, None, [tool], [action])
    args = out[0][1] if isinstance(out, list) and out else {}
    _check("[agent] draft message carries real values", json.dumps(args),
           [RAVI["mobile"], RAVI["pan"]])
    time.sleep(PAUSE)


def test_zz_report(outbound):
    hosts: dict[str, int] = {}
    for h, _ in outbound:
        hosts[h] = hosts.get(h, 0) + 1
    model = os.getenv("LLM_MODEL")
    print(f"\n\n=== Model: {model} | embedder: {os.getenv('EMBEDDING_BACKEND')} "
          f"{os.getenv('EMBEDDING_MODEL')} | rerank: {os.getenv('RETRIEVAL_RERANK_ENABLED')} ===")
    print("=== Outbound calls by host ===")
    for h, n in sorted(hosts.items()):
        print(f"  {h}: {n}")
    print(f"=== Raw identifiers found in outbound calls: {len(_leaks(outbound))} ===")
    print("=== Scenarios ===")
    for q, status, answer in REPORT:
        print(f"- {q}\n    {status}\n    {answer!r}")
